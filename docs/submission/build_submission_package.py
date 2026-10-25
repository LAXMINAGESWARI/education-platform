"""Build the CurioEd group submission package (DOCX + HTML/PDF).

The PDF is a print of the HTML. The DOCX is the editable source of the same
blocks, so the two stay in step. Run from anywhere:

    python3 docs/submission/build_submission_package.py
"""

from __future__ import annotations

import html
import os
import signal
import subprocess
import time
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
CHROME = "google-chrome"

NAVY = RGBColor(0x1E, 0x3A, 0x5F)
INK = RGBColor(0x1A, 0x23, 0x32)
MUTED = RGBColor(0x4A, 0x5A, 0x6A)

MEMBERS = [
    "Aushak Anuj Dendukuri",
    "Pranit Daggupati",
    "Suparna Dhumale",
    "Rahul Vamshikanth Gunda",
    "Arathe Sankar",
]

REPO = "https://github.com/aushakanuj/education-platform"
DATE = "25 October 2026"
DEADLINE = "25 October 2026, 11:55 PM"


def _p(text: str) -> tuple:
    return ("p", text)


def _h1(text: str) -> tuple:
    return ("h1", text)


def _h2(text: str) -> tuple:
    return ("h2", text)


def _ul(items: list[str]) -> tuple:
    return ("ul", items)


def _ol(items: list[str]) -> tuple:
    return ("ol", items)


def _table(headers: list[str], rows: list[list[str]]) -> tuple:
    return ("table", headers, rows)


def _note(text: str) -> tuple:
    return ("note", text)


def _break() -> tuple:
    return ("break",)


def _sig() -> tuple:
    return ("sig",)


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------

REPORT: list[tuple] = [
    ("cover", "Final Project Report", "CurioEd: a shared curriculum, a private record of learning, and an early-warning list a teacher can read."),
    _h1("1. Purpose"),
    _p(
        "CurioEd is a school platform for three jobs that usually get split across "
        "different tools. An administrator publishes one curriculum for a grade and "
        "a subject. Each enrolled student studies that same curriculum through a "
        "private learning record, sits the same quiz, and sees their own result. "
        "A teacher or an administrator can open a short list of students who may "
        "need a conversation, and every name on that list carries the reason it appeared."
    ),
    _p(
        "The repository is github.com/aushakanuj/education-platform. This report "
        "describes the system that is in that repository: what it does, how the "
        "early-warning rules were chosen, what the synthetic school is for, and "
        "what the system refuses to do. Paths below are from the repository root."
    ),
    _h1("2. Constraints the work had to respect"),
    _p(
        "On 19 August 2026 the group demonstrated the student, teacher, and "
        "administrator paths to mentor Peeyush. The minutes of that review "
        "(docs/mom/week2.md) record two directions. Analytics had to become a "
        "method someone can justify: which signals, which cut-offs, and why. "
        "AI help had to stay inside material the school had already approved, "
        "with a person in the loop, rather than inventing lesson facts on the fly. "
        "Sponsor feedback, relayed by Pranit from Lakshmi, asked for personal "
        "feedback after a quiz. Peeyush tied that feedback to question-level tags, "
        "so a weak area can be named instead of handing the student a single percentage."
    ),
    _p(
        "The sponsor did not provide historical pupil records. The generator "
        "module says this directly: CBT is providing no data, so the platform "
        "builds its own school. The early-warning design note draws the "
        "consequence. A learned model would need two or three cohorts of real "
        "outcomes the sponsor could not confirm. The signed direction was a "
        "rule a person can read, not a score a model emits. The design note "
        "also forbids presenting these rules as a prediction about a real child. "
        "Every figure in the demo comes from a generated school."
    ),
    _h1("3. Product model"),
    _p(
        "The product is two directories over one curriculum tree. The tree is "
        "Academic Period, Grade, Subject, Topic, Subtopic."
    ),
    _p(
        "SourceCurriculum belongs to the administrator. A subtopic holds published "
        "lesson versions and one common mastery quiz. Learner activity never "
        "overwrites that source."
    ),
    _p(
        "StudentLearningDirectory belongs to one student. It points at the "
        "published lesson and quiz. It stores that student's progress and "
        "attempts. It does not keep a second copy of the file. A student reaches "
        "it only with an active account, an active period, a grade enrollment, "
        "and a grade–subject enrollment, and only when the source is published."
    ),
    _p(
        "Teachers see evidence for classes they are assigned to teach. "
        "Administrators can read the institution directory without holding a "
        "student enrollment. A second institution sees nothing of the first. "
        "Roles in the schema are administrator, teacher, and student. Access "
        "is a JWT plus those roles and enrollments. The platform does not add "
        "a side door around them."
    ),
    _h1("4. What is running"),
    _p(
        "The proof-of-concept loop is: publish a lesson, study it, sit the common "
        "quiz, score it, and give staff an explainable flag. The pieces below "
        "are in the codebase and covered by the test suite. Items still marked "
        "as a design contract, rather than a shipped behaviour, are listed in "
        "section 9."
    ),
    _table(
        ["Area", "Where it lives", "State"],
        [
            [
                "Sign-in, refresh tokens, roles",
                "modules/auth",
                "Running. Demo school password is the fixture demo1234.",
            ],
            [
                "Periods, grades, subjects, enrollments, learning directory",
                "modules/academics",
                "Running. Students need active grade and grade–subject enrollments.",
            ],
            [
                "Lessons, progress, curriculum seed, PDF ingest",
                "modules/materials, workers",
                "Running. Approved markdown in docs/curriculum is the seed. PDF blobs stay on local disk.",
            ],
            [
                "Common quizzes and attempts",
                "modules/assessments",
                "Running. Default pass mark is 70%. Student APIs do not return answer keys.",
            ],
            [
                "student_360 register",
                "modules/insights",
                "Running. One read path for mastery, quizzes taken, and attendance.",
            ],
            [
                "Early-warning flags",
                "modules/at_risk, teacher and admin /at-risk pages",
                "Running. Staff only. Each flag names its drivers. Dismissal and viewing are audited.",
            ],
            [
                "Ask-the-data over the same flags",
                "modules/text_to_sql, at-risk agreement tests",
                "Running. A teacher question and the flag list are tested to return the same students.",
            ],
            [
                "Policy assistant",
                "modules/assistant",
                "Running for administrators. Pipeline is inject, validate, retrieve, summarize, over indexed documents.",
            ],
            [
                "Vector index",
                "modules/rag, Postgres pgvector",
                "Running. Ingest jobs are claimed with FOR UPDATE SKIP LOCKED.",
            ],
        ],
    ),
    _p(
        "Demo accounts for the synthetic Al Noor school, all with password "
        "demo1234: administrator fatima.almansouri@alnoor.school, teacher "
        "meera.krishnan@alnoor.school, student student1@alnoor.school (Adnan, "
        "Grade 6), and student97@alnoor.school (Aisha Rahman, Grade 8)."
    ),
    _h1("5. Early-warning method"),
    _h2("5.1 What a flag is"),
    _p(
        "A flag is a staff prompt to look at one student. It is not a grade, "
        "not a letter home, and not a label shown to the student. The design "
        "note docs/design/08-at-risk-early-warning.md separates it from the "
        "per-subtopic struggle signal a student may see on their own feedback "
        "page. Struggle says “this topic is weak.” At-risk says “a teacher "
        "should look at this child,” and it has to say why."
    ),
    _p(
        "Every input is read from student_360. Subject mastery and quiz counts "
        "are per student per subject offering. Attendance is per student for "
        "the whole period, because that is the grain the register actually "
        "stores. A rule that pretended attendance was per subject would be "
        "wrong on this schema. The engine is a pure function in "
        "modules/at_risk/engine.py. It does not query the database. The "
        "service loads signals and stores the flags the function returns."
    ),
    _h2("5.2 Rules"),
    _p(
        "Mastery and attendance are judged separately. They are never averaged "
        "into one blend. A blend would hide “fine everywhere except mathematics” "
        "inside a middling composite, which is the case the design note exists "
        "to catch."
    ),
    _ul(
        [
            "Mastery level. Flag the subject when mastery_percent is below 60.",
            "Mastery trend. When the student has at least two attempts before the latest three, compare the mean of those earlier attempts with the mean of the latest three. Flag when the later mean is more than 15 points lower.",
            "Attendance level. Flag when attendance_percent is below 80. There is no attendance trend in this version: the register stores a cumulative percentage, not a series.",
            "A subject flag never lists attendance as one of its drivers. An attendance flag names no subject.",
            "A flag with an empty driver list cannot be constructed. That check is in the type itself (requirement AR-1).",
        ]
    ),
    _p("Tiers follow from how many drivers fired, and from which kind:"),
    _table(
        ["Tier", "When it is used"],
        [
            ["urgent", "Two or more drivers on the same flag. In practice, low mastery and a sharp decline together."],
            ["attention", "The only driver is the trend. The latest level can still sit at or above 60."],
            ["monitor", "A single level driver: mastery below 60, or attendance below 80."],
        ],
    ),
    _p(
        "The cuts are closed on the safe side of the line. Mastery of exactly 60, "
        "a decline of exactly 15 points, and attendance of exactly 80 do not "
        "flag. Missing attendance produces no attendance flag and does not raise "
        "an error. Too little history to fill the trend window produces no trend "
        "driver. The level rule still applies."
    ),
    _h2("5.3 Why these numbers"),
    _p(
        "The design note refuses to invent a cut-off with no distribution behind "
        "it. The engine file records the calibration against the generated school:"
    ),
    _ul(
        [
            "Mean mastery on the synthetic school is recorded as 71.3%, standard deviation 8.5. One and a half standard deviations below that mean is about 58.5%. The level cut is 60, just above that. The same note records that this flags about 11.5% of student–subject rows.",
            "A decline of more than 15 points is recorded for about 2.9% of groups that have enough history. It is meant to be uncommon. A slide of a few points is ordinary noise.",
            "Attendance at 80 sits in a gap the same note describes: a cluster near 70%, then the next mass of students at 86% and above. The old on-screen placeholder of 75 was a demo constant and has been removed from the roster.",
        ]
    ),
    _p(
        "These constants belong to this generated school. Reusing them on another "
        "school, without looking at that school's own distribution, would repeat "
        "the mistake the design note was written to stop. A different rule, exam "
        "eligibility, still uses 75% attendance (a student below that line cannot "
        "sit the end-of-term exam). That rule is not the early-warning cut. The "
        "two numbers are supposed to differ, and the code names them differently "
        "so they do not drift back into one shared constant."
    ),
    _h2("5.4 A worked case: Aisha Rahman"),
    _p(
        "Aisha Rahman is student S-00097 in Grade 8. She is planted by the "
        "generator so the engine has a case it must get right. The engine test "
        "records her mathematics sequence as it was read from the generated "
        "database on 29 August 2026. Chronological scores are 74, 65, 56, 47, 38. "
        "Her subject mastery is 56%. The design note records her other subjects "
        "between 61% and 77%, and her attendance at 62%."
    ),
    _p("Mathematics, level. 56 is below 60, so the mastery_percent driver fires. On its own that would be a monitor flag."),
    _p(
        "Mathematics, trend. The engine reads scores most-recent first: 38, 47, "
        "56, 65, 74. The latest three average (38 + 47 + 56) / 3 = 47. The "
        "earlier two average (65 + 74) / 2 = 69.5. The decline is 69.5 − 47 = "
        "22.5 points, which is past 15. The mastery_trend driver fires, and the "
        "test locks that 22.5."
    ),
    _p(
        "Two drivers on the mathematics flag make the tier urgent. The flag "
        "names the mathematics offering. A strong subject for the same student "
        "produces no flag. The test for that pair asserts a single flag, and "
        "that its offering is mathematics."
    ),
    _p(
        "Attendance. 62 is below 80, so a second flag fires with only an "
        "attendance_percent driver, tier monitor, and no subject id. Attendance "
        "is not copied onto the mathematics flag. The inverse is also tested: "
        "poor attendance does not leak into a subject driver. Because the "
        "attendance flag names no subject, the visibility rule sends it to "
        "administrators. The mathematics flag is visible to the teacher assigned "
        "to that offering, and to administrators."
    ),
    _h2("5.5 The other planted shapes"),
    _ul(
        [
            "Attendance without weak marks. Every 17th student is generated at a 70% attendance rate with marks that stay clear of the mastery cut. The engine returns one flag, attendance only, tier monitor. A fine subject contributes nothing.",
            "Decline while the level still looks acceptable. A trend-only flag is tier attention. A level-only rule would have missed this student. That is why both conditions exist.",
            "A section gap. Two sections of the same subject are built so that one topic really does differ. The individual rule still emits ordinary per-student flags. It does not invent a cohort object that would hide a teaching or content gap inside a pile of identical personal flags. The design note leaves the cohort story itself out of scope, and requires the engine not to paper over it.",
        ]
    ),
    _h2("5.6 What staff are allowed to do with a flag"),
    _p("The design note binds five behaviours. The API tests exercise them."),
    _table(
        ["Id", "Rule"],
        [
            ["AR-1", "No flag without at least one named driver (metric, value, comparison, window)."],
            ["AR-2", "A flag never sends a message to a student, parent, or guardian."],
            ["AR-3", "A flag never writes a grade or any other academic record."],
            ["AR-4", "Dismissal is optional-note, mandatory-audit: who, when, which flag."],
            ["AR-5", "Opening the flag list is an audited read, scoped like any other student_360 read."],
        ],
    ),
    _p(
        "Students receive 403 on the flag list, on dismiss, and on recompute. "
        "A teacher sees flags only for subjects they teach. An administrator sees "
        "more, including attendance-only flags, which a subject teacher does not "
        "see. An administrator of another institution sees an empty list. "
        "Dismissing a flag outside one's scope returns 404, the same shape as a "
        "missing flag, so the response does not reveal that the row exists. "
        "After dismissal the flag leaves the active list, and both the view and "
        "the dismissal write an audit event."
    ),
    _h1("6. What the tests fix in place"),
    _p(
        "Unit tests in backend/tests/test_at_risk_engine.py cover the boundary "
        "values, the empty-driver rejection, Aisha's mathematics sequence, the "
        "attendance-only shape, and the refusal to merge two weak subjects into "
        "one composite. API tests in backend/tests/test_at_risk_api.py cover "
        "role, tenant, dismiss, and audit."
    ),
    _p(
        "A separate agreement fixture, not the 240-student school, checks that "
        "the staff question-answering path and the flag endpoint name the same "
        "students. In that fixture the teacher list has 9 students, all in "
        "mathematics: eight planted low-mastery students and one trend-only "
        "student. A legacy query that also counted low attendance returned 12. "
        "The four attendance-only students are on the administrator list and "
        "absent from the teacher list. The test exists so a later prompt cannot "
        "quietly resurrect the legacy count."
    ),
    _p(
        "Continuous integration on pull requests runs Ruff format, Ruff lint, "
        "mypy on the backend, Alembic upgrade, and pytest with an 80% coverage "
        "gate, and the frontend Vitest suite in parallel. Live calls to a paid "
        "language model are opt-in and are not part of that job."
    ),
    _h1("7. Data"),
    _p(
        "The one-page Data Note in this package is the data statement. In short: "
        "there is no extract of real pupils. The default school is Al Noor "
        "International, Grades 6 to 10, two sections per grade, 24 students per "
        "section (240 students), five subjects, five quizzes per subject, ten "
        "term weeks, seed 20260812. The same command rebuilds the same school. "
        "Names, marks, and attendance are authored so that the planted cases in "
        "section 5 have a determinate answer. Nothing in that school is evidence "
        "about a living child."
    ),
    _h1("8. Grounded assistance"),
    _p(
        "The administrator policy assistant is a fixed pipeline: check the "
        "question for injection, validate it, retrieve chunks from the school's "
        "indexed documents, then summarize. Retrieval is pgvector over those "
        "chunks. If no OpenRouter key is configured, a heuristic injection check "
        "and a stub summary still run, so the path is exercisable without a "
        "paid model. The key is an environment variable. It is not part of this "
        "submission."
    ),
    _p(
        "Student-facing answers are a narrower contract. Quiz APIs do not return "
        "the answer key. Scoring happens on the server. The design for lesson "
        "generation is the same idea the mentor asked for in August: a draft "
        "has to come from approved sources, and a person publishes it. A "
        "student grounded tutor, sitting on top of the same index, is still a "
        "later piece of work. The policy assistant is the part that runs today."
    ),
    _h1("9. Limits"),
    _ul(
        [
            "The school is synthetic. Thresholds tuned on it show that the mechanism distinguishes the planted cases. They do not show that the mechanism predicts real outcomes.",
            "The engine is rules, by decision. A learned scorer is out of scope, not a backlog item waiting on more of the same synthetic rows.",
            "Four-pillar evaluation snapshots (marks, anonymised peer band, weak outcomes, trend) are specified in docs/design/06-analytics-and-comparison-insights.md. Peer bands require at least five scored students and 80% completion before a comparison is shown. Stored snapshots of that full object are not the shipped record. What ships today is attempt scoring, feedback from those attempts, student_360, and the early-warning flags.",
            "Some teacher screens outside the at-risk list are still fed by fixtures. The at-risk pages call the live API.",
            "There is no parent role, and a flag cannot notify a family.",
            "Lesson files for the proof of concept live on the application disk. Shared object storage is not in this version.",
            "A native mobile client is a target architecture, not a shipped client. The web client is the one that runs.",
        ]
    ),
    _h1("10. Reproducing the analysis"),
    _ol(
        [
            "Start Postgres with pgvector from the repository root: docker compose up -d postgres.",
            "From backend/, apply migrations: uv run alembic upgrade head.",
            "Rebuild the school: uv run python -m education_platform.modules.synthetic.cli. The default seed is 20260812.",
            "The engine does not need the database. From backend/: uv run pytest tests/test_at_risk_engine.py tests/test_at_risk_api.py tests/test_at_risk_assistant_agreement.py.",
            "Accounts, curls, and the web client are in the repository README.",
        ]
    ),
    _p(
        "The Code and Analysis Repository note in this package is the short map "
        "of those paths and the link to submit for that line of the LMS package."
    ),
    _h1("11. People"),
    _p(
        "Mentor for the 19 August 2026 review: Peeyush. Sponsor contact named in "
        "those minutes: Lakshmi. The minutes also list tanmay among the "
        "attendees. This repository has no commit under that name and no "
        "surname for that person, so tanmay is not on the signature block. Add "
        "a row on the declaration if tanmay is a submitting member."
    ),
    _p(
        "Human commit counts below exclude automated-agent commits. They are a "
        "trace of who touched the tree, not a split of ownership. Everyone "
        "listed signs one group declaration."
    ),
    _table(
        ["Name", "Trace in the repository"],
        [
            ["Aushak Anuj Dendukuri", "Repository owner (github.com/aushakanuj). Largest set of human commits on the platform."],
            ["Pranit Daggupati", "Platform commits. Minutes record him carrying sponsor feedback into the 19 August review."],
            ["Suparna Dhumale", "Platform commits. Minutes credit the Student 360 view."],
            ["Arathe Sankar", "Struggle-flag readout and the student feedback presentation (TEAM-READOUT.md, September 2026)."],
            ["Rahul Vamshikanth Gunda", "Commits on the repository under Rahul G."],
        ],
    ),
    _h1("12. This package"),
    _p(
        "The LMS notice asks for one submission per group, due 25 October 2026 "
        "at 11:55 PM. The notice does not name a time zone. The group's recorded "
        "meeting used India Standard Time. The package beside this report is:"
    ),
    _ul(
        [
            "This report, as PDF and as DOCX.",
            "Assignment Submission Form.",
            "Data Note, one page.",
            "Originality and Confidentiality Declaration, with signature lines left blank for the group.",
            "Plain-Language Summary.",
            "Code and Analysis Repository note, pointing at the GitHub repository.",
            "A pre-upload checklist. The notice points at Section 9 of an Instructions and Best Practices guide. That guide is not in the repository, so the checklist covers the items the notice itself named, and it says so.",
        ]
    ),
    _p(
        "The narrative documents were drafted with Cursor from the repository's "
        "code, tests, and design notes. The declaration states that. Signing "
        "means the group has checked the statements against the code."
    ),
]


FORM: list[tuple] = [
    (
        "cover",
        "Assignment Submission Form",
        "One Submission Package for the group. Due 25 October 2026 at 11:55 PM.",
    ),
    _h1("1. Submission"),
    _table(
        ["Field", "Entry"],
        [
            ["Project title", "CurioEd: a shared curriculum, a private learning record, and an explainable early-warning list"],
            ["Product name", "CurioEd"],
            ["Repository", REPO],
            ["Deadline on the LMS notice", f"{DEADLINE}. The notice does not name a time zone. Group meetings were recorded in India Standard Time."],
            ["How many uploads", "One submission for the whole group."],
            ["Course or group number", "Not recorded in the repository. Write it here before upload: ________________"],
            ["Mentor named in the 19 August 2026 minutes", "Peeyush"],
            ["Sponsor contact named in those minutes", "Lakshmi"],
            ["Document date", DATE],
        ],
    ),
    _h1("2. Package inventory"),
    _p(
        "Tick the paper copy or the LMS file list against this inventory. "
        "The PDF and the DOCX of each narrative document carry the same text."
    ),
    _table(
        ["LMS item", "File in this package", "In the package"],
        [
            ["Final Project Report (PDF)", "01-Final-Project-Report.pdf", "Yes"],
            ["Editable source file (.docx)", "01-Final-Project-Report.docx", "Yes. Same text as the PDF."],
            ["Assignment Submission Form", "02-Assignment-Submission-Form.pdf and .docx", "Yes. This document."],
            ["Code and Analysis Repository", REPO + " — see 06-Code-and-Analysis-Repository.pdf", "Link. Analysis is Python modules and tests, not a notebook bundle."],
            ["Data Note (maximum 1 page)", "03-Data-Note.pdf and .docx", "Yes. Confirm the PDF is a single page before upload."],
            ["Signed Originality and Confidentiality Declaration", "04-Originality-and-Confidentiality-Declaration.pdf and .docx", "Text is complete. Signature lines are blank until the group signs."],
            ["Plain-Language Summary", "05-Plain-Language-Summary.pdf and .docx", "Yes."],
            ["Pre-upload checklist", "09-Pre-Upload-Checklist.pdf and .docx", "Working checklist for the items the LMS notice named. See that sheet for the Section 9 gap."],
        ],
    ),
    _h1("3. Members"),
    _p(
        "These five names are the people with a full name in the repository "
        "history or the 19 August 2026 minutes. tanmay is named in the minutes "
        "with no surname and no commit. Add a row if tanmay is a submitting "
        "member. Each person signs the declaration, not this form, unless the "
        "course asks for both."
    ),
    _table(
        ["Name", "How they appear in the repository"],
        [
            ["Aushak Anuj Dendukuri", "GitHub owner of aushakanuj/education-platform. Platform commits."],
            ["Pranit Daggupati", "Platform commits. Sponsor feedback at the 19 August review."],
            ["Suparna Dhumale", "Platform commits. Student 360 view, per the minutes."],
            ["Rahul Vamshikanth Gunda", "Commits as Rahul G."],
            ["Arathe Sankar", "Struggle-flag readout and student feedback work, September 2026."],
        ],
    ),
    _h1("4. Statements the uploading member is making"),
    _p("The person who uploads the single group package confirms, by uploading:"),
    _ol(
        [
            "This is the one submission for the group.",
            "The report PDF and the report DOCX are the same text.",
            "The Data Note PDF is one page.",
            "The repository link above is the code and analysis for this project, and it matches the report.",
            "The declaration has been signed by the submitting members before upload. An unsigned declaration does not meet the LMS line that asks for a signed declaration.",
            "The upload contains no API keys, no production passwords, and no records of real pupils. The Al Noor school is synthetic. The demo password demo1234 is a local fixture, documented in the repository README.",
        ]
    ),
    _h1("5. Uploading member"),
    _p("Name: ________________________________"),
    _p("Signature: ________________________________"),
    _p("Date: ________________________________"),
    _p("Time zone of the clock used: ________________________________"),
]


DATA_NOTE: list[tuple] = [
    (
        "cover",
        "Data Note",
        "One page. The school behind every mark in CurioEd is generated. No row is a real pupil.",
    ),
    _p(
        "CurioEd did not receive an extract of real students. The generator module "
        "states that CBT provided no pupil records, and the early-warning design "
        "note records that the sponsor could not confirm historical outcome "
        "cohorts. Marks, attendance, names, and classes are produced by "
        "education_platform.modules.synthetic from a fixed seed. The same command "
        "rebuilds the same school. This note is the data statement for that school "
        "and for the rules that read it."
    ),
    _table(
        ["Item", "Value"],
        [
            ["Institution name in the demo", "Al Noor International"],
            ["Grades", "6, 7, 8, 9, 10"],
            ["Sections and students", "2 sections per grade, 24 students per section, 240 students"],
            ["Subjects", "Mathematics, Science, English, Arabic, Social Studies"],
            ["Quizzes and term", "5 quizzes per subject, 10 term weeks"],
            ["Seed", "20260812"],
            ["Demo password", "demo1234, a local fixture, not a production secret"],
            ["Generator", "backend/: uv run python -m education_platform.modules.synthetic.cli"],
        ],
    ),
    _p(
        "Three shapes are planted on purpose, because a uniform random school "
        "makes every rule look the same. Aisha Rahman (S-00097, Grade 8) declines "
        "in mathematics: scores 74, 65, 56, 47, 38, mastery 56%, attendance rate "
        "62%, other subjects recorded in the design note between 61% and 77%. "
        "Every 17th student has attendance at 70% and marks that are not low. "
        "Two sections of one subject differ on a mathematics topic, so a section "
        "gap is real and is not collapsed into a single cohort flag."
    ),
    _p(
        "Rules, recorded with their calibration in modules/at_risk/engine.py: "
        "mastery below 60% (synthetic mean mastery 71.3%, SD 8.5; about 11.5% of "
        "student–subject rows); a drop of more than 15 points from earlier "
        "attempts to the latest three (about 2.9% of groups with history); "
        "attendance below 80% (the gap between a 70% cluster and the next mass "
        "at 86%). Exact hits on 60, 15, and 80 do not flag. A subject flag and "
        "an attendance flag are separate objects. These cuts describe this "
        "generated school only."
    ),
    _p(
        "The data can show that those planted cases fire the intended drivers, "
        "and that a teacher and an administrator do not see the same rows. It "
        "cannot support a prediction about a real child, a trained model, or a "
        "comparison with a live school. Do not treat Al Noor International as "
        "a real institution. Policy: docs/design/08-at-risk-early-warning.md."
    ),
]


DECLARATION: list[tuple] = [
    (
        "cover",
        "Originality and Confidentiality Declaration",
        "The text is the group's statement. It takes effect when the submitting members sign.",
    ),
    _h1("1. Originality"),
    _p(
        "By signing, each person confirms that their contribution to CurioEd, "
        "and to this Submission Package, is their own work or the joint work of "
        "this group, except for the third-party components named below. The "
        "group is not submitting this project for assessment in another course "
        "under a different description of authorship. Where a person reused "
        "their own earlier notes, they can say so in the margin before signing."
    ),
    _h1("2. Third-party software"),
    _p(
        "The platform is built on software the group did not write. That "
        "software stays under its own licences. The lockfiles are the list: "
        "backend/uv.lock and frontend/package.json (and the frontend lockfile). "
        "Major pieces include Python, FastAPI, SQLAlchemy, PostgreSQL with "
        "pgvector, LangGraph, React, and Vite. The group claims the application "
        "code, the curriculum markdown it authored, the design notes, the tests, "
        "and this narrative. The group does not claim those libraries."
    ),
    _h1("3. No real pupil records"),
    _p(
        "The school called Al Noor International is synthetic. Names are drawn "
        "from lists in the generator. Marks and attendance are computed from "
        "seed 20260812. Nothing in the database is a record of a living child, "
        "and this package does not attach one. Demo accounts and the password "
        "demo1234 are local fixtures, documented in the repository README, for "
        "a database the group generates itself."
    ),
    _h1("4. Confidentiality"),
    _p(
        "Signing members will keep the following out of the upload and out of "
        "any public copy they do not intend to publish:"
    ),
    _ul(
        [
            "API keys, including any OPENROUTER_API_KEY. Keys live in uncommitted environment files.",
            "Passwords other than the documented demo fixture.",
            "Any pupil, parent, or staff record from a real school, should one ever be received. None is in this package.",
            "Sponsor material the sponsor has not agreed to include. The minutes name Lakshmi's feedback in summary. They do not attach a sponsor document.",
        ]
    ),
    _p(
        "Results from the synthetic school will be described as results from "
        "the synthetic school. They will not be presented as measurements of "
        "a real cohort."
    ),
    _h1("5. Drafting assistance"),
    _p(
        "The narrative documents in this package (the report, this declaration, "
        "the form, the data note, the plain-language summary, the repository "
        "note, and the checklist) were drafted with Cursor, an AI coding "
        "assistant, from the group's repository: code, tests, and design notes. "
        "No AI system is an author. Each signatory is responsible for the "
        "accuracy of the text they sign. If a statement disagrees with the "
        "code, the code is the record and the text should be corrected before "
        "upload."
    ),
    _h1("6. One group submission"),
    _p(
        "The group submits one package. The uploading member is named on the "
        "Assignment Submission Form. Signing here means the person has had the "
        "chance to read the report, the data note, and this declaration, and "
        "agrees to be named as a member of the group that submits them."
    ),
    _h1("7. Signatures"),
    _p(
        "Sign in ink or with a proper digital signature. A typed name with no "
        "signature is not a signature. Leave unused rows blank rather than "
        "signing for someone else. Add a row for any submitting member who is "
        "not listed. The minutes name tanmay with no surname; add that row if "
        "they are a submitting member."
    ),
    _sig(),
    _p("Document prepared " + DATE + "."),
]


PLAIN: list[tuple] = [
    (
        "cover",
        "Plain-Language Summary",
        "What CurioEd is, in ordinary words.",
    ),
    _p(
        "CurioEd is a website for a school. It has three kinds of login: an "
        "administrator, a teacher, and a student."
    ),
    _p(
        "The administrator publishes the lessons and the quizzes. There is one "
        "copy of each lesson for the class. Students do not each get a private "
        "rewrite of the file. They each get a private record of what they have "
        "done: which lessons they opened, and how they did on the quiz."
    ),
    _p(
        "A student can sign in, open their subjects, read a lesson, and take "
        "the quiz for that topic. The passing mark used by the quizzes is 70%. "
        "The site scores the quiz on the server. It does not send the student "
        "the answer key."
    ),
    _p(
        "A teacher can open a list of students who may need a conversation. "
        "Each row says why it is there. The reasons are simple. Marks in one "
        "subject are under 60%. Or the recent quizzes have fallen by more than "
        "15 points compared with earlier ones. Or attendance is under 80%. If "
        "two of those subject reasons are true at once, the row is marked "
        "urgent. If marks have fallen but are not yet under 60%, the row is "
        "marked attention. A single low score, or low attendance on its own, "
        "is marked monitor."
    ),
    _p(
        "A low mark in mathematics does not get described as “this student is "
        "weak at everything.” The other subjects stay off the list when their "
        "marks are fine. Low attendance is its own row, and a subject teacher "
        "does not see that row. An administrator does."
    ),
    _p(
        "The student is not shown an “at risk” label. The list is for staff. "
        "The site does not email a parent, and it does not change a grade "
        "because a name appeared. A teacher can dismiss a row. The site keeps "
        "a record of who dismissed it."
    ),
    _p(
        "The school in the demo is invented. It is called Al Noor International. "
        "It has 240 generated students in Grades 6 to 10. One of them, Aisha "
        "Rahman, is written so that her mathematics marks fall across the term "
        "(74, then 65, 56, 47, 38) while her attendance sits at 62%. The list "
        "should show her mathematics row as urgent, and her attendance as a "
        "separate reason. The numbers were chosen so the demo has a clear case, "
        "not because they describe a real child. They should not be quoted as "
        "a finding about a real school."
    ),
    _p(
        "There is also a chat for administrators that answers from documents "
        "the school has uploaded, after a check that the question is not trying "
        "to hijack the tool. A tutor that talks to students in the same way is "
        "not part of this version."
    ),
    _p(
        "The code is at " + REPO + ". The longer explanation, including the "
        "arithmetic for Aisha's quizzes, is in the Final Project Report. The "
        "Data Note is the one-page description of the generated school."
    ),
]


REPO_NOTE: list[tuple] = [
    (
        "cover",
        "Code and Analysis Repository",
        "The LMS item asks for a link or a zip of notebooks. This project submits the link. The analysis is Python, tested with pytest.",
    ),
    _h1("1. Link"),
    _p(REPO),
    _p(
        "Owner on GitHub: aushakanuj. Default branch: main. This Submission "
        "Package is added on a branch and opened as a pull request so the "
        "files can be reviewed; the link above is the project the report describes."
    ),
    _h1("2. What “analysis” means here"),
    _p(
        "There is no notebook bundle. The analysis is a small rule engine with "
        "a deterministic generated school behind it, plus tests that lock the "
        "worked cases in the report."
    ),
    _table(
        ["Role", "Path"],
        [
            ["Generated school", "backend/src/education_platform/modules/synthetic/"],
            ["Rebuild command", "From backend/: uv run python -m education_platform.modules.synthetic.cli"],
            ["Seed and size", "SchoolSpec in generator.py. Seed 20260812. 240 students by default."],
            ["Early-warning rules", "backend/src/education_platform/modules/at_risk/engine.py"],
            ["Loading, storage, permissions", "backend/src/education_platform/modules/at_risk/service.py and router.py"],
            ["student_360 register", "backend/src/education_platform/modules/insights/"],
            ["Policy the rules implement", "docs/design/08-at-risk-early-warning.md"],
            ["Engine tests", "backend/tests/test_at_risk_engine.py"],
            ["API, role, and audit tests", "backend/tests/test_at_risk_api.py"],
            ["Teacher list matches the assistant", "backend/tests/test_at_risk_assistant_agreement.py"],
            ["Schema the questions are written against", "backend/src/education_platform/modules/text_to_sql/config/schema_catalog.md"],
            ["How to run the whole app", "README.md at the repository root"],
        ],
    ),
    _h1("3. Run the analysis"),
    _ol(
        [
            "Install uv, Docker, and Node 18 or newer.",
            "From the repository root: docker compose up -d postgres. Wait until it is healthy.",
            "From backend/: uv run alembic upgrade head",
            "From backend/: uv run python -m education_platform.modules.synthetic.cli",
            "From backend/: uv run pytest tests/test_at_risk_engine.py tests/test_at_risk_api.py tests/test_at_risk_assistant_agreement.py",
            "Optional, the app itself: uv run uvicorn education_platform.main:app --reload --host 127.0.0.1 --port 8000, and from frontend/ npm run dev after copying .env.example to .env.",
        ]
    ),
    _p(
        "The engine tests include cases that do not need a database. The API "
        "and agreement tests need the migrated Postgres from step 2. CI on "
        "GitHub runs the full backend suite against a pgvector container."
    ),
    _h1("4. What not to upload beside the link"),
    _ul(
        [
            "backend/.env if it contains an API key.",
            "The local upload directory of PDF blobs.",
            "A dump of a database, even the synthetic one. The seed rebuilds it.",
            "node_modules or the Python virtual environment.",
        ]
    ),
    _p(
        "A zip of the whole working tree would be a worse copy than the link: "
        "it goes stale, and it is easy to sweep a secret into it. If the LMS "
        "form will not accept a link without a file, upload this note (PDF and "
        "DOCX) and paste the link in the form field."
    ),
]


CHECKLIST: list[tuple] = [
    (
        "cover",
        "Pre-upload Checklist",
        "Use this before the single group upload. Due 25 October 2026 at 11:55 PM.",
    ),
    _h1("1. About Section 9"),
    _p(
        "The LMS notice says to use the checklist in Section 9 of the "
        "Instructions and Best Practices guide, and that Section 1 of that "
        "guide gives the format for each item. That guide is not in the "
        "CurioEd repository and was not attached to the request to prepare "
        "this package. This sheet is the working checklist for the items the "
        "notice itself printed. If the guide adds a format rule this sheet "
        "does not mention, follow the guide and adjust the file before upload."
    ),
    _h1("2. Files"),
    _table(
        ["Check", "What to confirm", "Status while preparing"],
        [
            ["Report PDF", "01-Final-Project-Report.pdf opens and matches the DOCX.", "Prepared"],
            ["Report DOCX", "01-Final-Project-Report.docx is the editable source.", "Prepared"],
            ["Form", "02-Assignment-Submission-Form is filled. Course or group number is still blank if the repository never recorded one.", "Write the group number"],
            ["Repository", "The link https://github.com/aushakanuj/education-platform opens for the marker.", "Confirm access"],
            ["Data Note", "03-Data-Note.pdf is exactly one page.", "Check the page count"],
            ["Declaration", "04-Originality-and-Confidentiality-Declaration is signed by each submitting member.", "Not signed yet"],
            ["Plain-language summary", "05-Plain-Language-Summary is in the upload.", "Prepared"],
            ["One upload", "A single member uploads the whole package once.", "Choose the uploader"],
        ],
    ),
    _h1("3. Content checks"),
    _ul(
        [
            "The report's worked example for Aisha Rahman matches the engine test: scores 74, 65, 56, 47, 38, decline 22.5 points, tier urgent on mathematics, attendance separate.",
            "The Data Note does not call the synthetic school a real cohort.",
            "No OPENROUTER_API_KEY, no .env file, and no database dump is in the upload.",
            "The demo password appears only as the documented fixture demo1234.",
            "Answer keys are not pasted into the report.",
            "The declaration's drafting-assistance paragraph is still true.",
            "tanmay is added to the signature block if they are a submitting member.",
        ]
    ),
    _h1("4. Deadline"),
    _p(
        "The notice says the package is due on 25 October 2026 at 11:55 PM, "
        "one submission per group, with late work handled under the course "
        "policy. The notice does not print a time zone. The group's meeting "
        "minutes use India Standard Time. Upload against the clock the LMS "
        "uses, and write that zone on the form."
    ),
]


DOCUMENTS: list[tuple[str, str, list[tuple]]] = [
    ("01-Final-Project-Report", "Final Project Report", REPORT),
    ("02-Assignment-Submission-Form", "Assignment Submission Form", FORM),
    ("03-Data-Note", "Data Note", DATA_NOTE),
    (
        "04-Originality-and-Confidentiality-Declaration",
        "Originality and Confidentiality Declaration",
        DECLARATION,
    ),
    ("05-Plain-Language-Summary", "Plain-Language Summary", PLAIN),
    ("06-Code-and-Analysis-Repository", "Code and Analysis Repository", REPO_NOTE),
    ("09-Pre-Upload-Checklist", "Pre-upload Checklist", CHECKLIST),
]


def _esc(text: str) -> str:
    return html.escape(text, quote=True)


def _runs(text: str) -> str:
    """Minimal inline code formatting with backticks."""
    parts = text.split("`")
    out = []
    for index, part in enumerate(parts):
        escaped = _esc(part)
        if index % 2 == 1:
            out.append(f"<code>{escaped}</code>")
        else:
            out.append(escaped)
    return "".join(out)


def render_html(title: str, blocks: list[tuple], *, compact: bool) -> str:
    body: list[str] = []
    for block in blocks:
        kind = block[0]
        if kind == "cover":
            _, heading, deck = block
            body.append(
                "<header class='cover'>"
                f"<p class='kicker'>CurioEd · group submission · { _esc(DATE) }</p>"
                f"<h1>{_esc(heading)}</h1>"
                f"<p class='deck'>{_esc(deck)}</p>"
                f"<p class='meta'>{_esc(REPO)}</p>"
                "</header>"
            )
        elif kind == "h1":
            body.append(f"<h2>{_esc(block[1])}</h2>")
        elif kind == "h2":
            body.append(f"<h3>{_esc(block[1])}</h3>")
        elif kind == "p":
            body.append(f"<p>{_runs(block[1])}</p>")
        elif kind == "note":
            body.append(f"<aside>{_runs(block[1])}</aside>")
        elif kind == "ul":
            items = "".join(f"<li>{_runs(item)}</li>" for item in block[1])
            body.append(f"<ul>{items}</ul>")
        elif kind == "ol":
            items = "".join(f"<li>{_runs(item)}</li>" for item in block[1])
            body.append(f"<ol>{items}</ol>")
        elif kind == "table":
            _, headers, rows = block
            head = "".join(f"<th>{_esc(cell)}</th>" for cell in headers)
            body_rows = []
            for row in rows:
                cells = "".join(f"<td>{_esc(cell)}</td>" for cell in row)
                body_rows.append(f"<tr>{cells}</tr>")
            body.append(
                "<table><thead><tr>"
                + head
                + "</tr></thead><tbody>"
                + "".join(body_rows)
                + "</tbody></table>"
            )
        elif kind == "sig":
            rows = []
            for name in MEMBERS:
                rows.append(
                    "<tr>"
                    f"<td>{_esc(name)}</td>"
                    "<td class='sign'></td>"
                    "<td class='sign'></td>"
                    "</tr>"
                )
            rows.append(
                "<tr><td>Additional member (name in full)</td>"
                "<td class='sign'></td><td class='sign'></td></tr>"
            )
            body.append(
                "<table class='sig'><thead><tr>"
                "<th>Name</th><th>Signature</th><th>Date</th>"
                "</tr></thead><tbody>"
                + "".join(rows)
                + "</tbody></table>"
            )
        elif kind == "break":
            body.append("<div class='break'></div>")
        else:
            raise ValueError(kind)

    density = "compact" if compact else "regular"
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8"/>
<title>{_esc(title)}</title>
<style>
  @page {{
    size: A4;
    margin: 16mm 15mm 18mm 15mm;
    @bottom-left {{
      content: "CurioEd  ·  {_esc(title)}";
      font-family: "Inter", "Noto Sans", sans-serif;
      font-size: 8.5pt;
      color: #5c6b7a;
    }}
    @bottom-right {{
      content: counter(page);
      font-family: "Inter", "Noto Sans", sans-serif;
      font-size: 8.5pt;
      color: #5c6b7a;
    }}
  }}
  * {{ box-sizing: border-box; }}
  html, body {{ margin: 0; padding: 0; }}
  body {{
    font-family: "Noto Serif", "Times New Roman", serif;
    color: #1a2332;
    font-size: 11pt;
    line-height: 1.45;
  }}
  body.compact {{
    font-size: 10pt;
    line-height: 1.32;
  }}
  .cover {{
    border-top: 8px solid #1e3a5f;
    padding-top: 14px;
    margin-bottom: 18px;
  }}
  .kicker {{
    font-family: "Inter", "Noto Sans", sans-serif;
    font-size: 9pt;
    letter-spacing: 0.04em;
    text-transform: uppercase;
    color: #5c6b7a;
    margin: 0 0 8px;
  }}
  h1 {{
    font-family: "Inter", "Noto Sans", sans-serif;
    font-weight: 650;
    font-size: 22pt;
    line-height: 1.15;
    color: #1e3a5f;
    margin: 0 0 10px;
  }}
  .deck {{
    font-size: 12pt;
    margin: 0 0 8px;
  }}
  body.compact .deck {{ font-size: 10.5pt; }}
  .meta {{
    font-family: "Inter", "Noto Sans", sans-serif;
    font-size: 9pt;
    color: #1e3a5f;
    margin: 0;
  }}
  h2 {{
    font-family: "Inter", "Noto Sans", sans-serif;
    font-size: 13.5pt;
    color: #1e3a5f;
    margin: 18px 0 6px;
    line-height: 1.25;
    page-break-after: avoid;
  }}
  h3 {{
    font-family: "Inter", "Noto Sans", sans-serif;
    font-size: 11.5pt;
    color: #1e3a5f;
    margin: 12px 0 4px;
    page-break-after: avoid;
  }}
  body.compact h2 {{ margin-top: 11px; font-size: 12pt; }}
  body.compact h3 {{ margin-top: 8px; }}
  p {{ margin: 0 0 8px; orphans: 3; widows: 3; }}
  ul, ol {{ margin: 0 0 8px; padding-left: 1.2em; }}
  li {{ margin: 0 0 3px; }}
  code {{
    font-family: "JetBrains Mono", ui-monospace, monospace;
    font-size: 0.86em;
  }}
  table {{
    width: 100%;
    border-collapse: collapse;
    margin: 6px 0 10px;
    page-break-inside: avoid;
    font-family: "Inter", "Noto Sans", sans-serif;
    font-size: 8.7pt;
    line-height: 1.3;
  }}
  body.compact table {{ font-size: 8pt; margin: 4px 0 6px; }}
  th {{
    background: #1e3a5f;
    color: white;
    text-align: left;
    font-weight: 600;
    padding: 5px 6px;
    vertical-align: top;
  }}
  td {{
    border-bottom: 1px solid #d5dde6;
    padding: 4px 6px;
    vertical-align: top;
  }}
  tr:nth-child(even) td {{ background: #f4f7fb; }}
  table.sig td {{ height: 46px; }}
  table.sig td.sign {{ background: white; }}
  aside {{
    border-left: 3px solid #1e3a5f;
    padding: 4px 10px;
    margin: 0 0 8px;
    color: #1a2332;
  }}
  .break {{ page-break-before: always; }}
</style>
</head>
<body class="{density}">
{''.join(body)}
</body>
</html>
"""


def _set_run_font(run, name: str, size_pt: float, *, bold: bool = False, color: RGBColor | None = None) -> None:
    run.bold = bold
    run.font.name = name
    run.font.size = Pt(size_pt)
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _shade(cell, fill: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_border(cell) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D5DDE6")
        tc_borders.append(element)
    tc_pr.append(tc_borders)


def _paragraph_text(paragraph, text: str, font: str, size: float, *, bold: bool = False, color: RGBColor | None = None) -> None:
    # Backtick spans become a monospace run.
    parts = text.split("`")
    if len(parts) == 1:
        run = paragraph.add_run(text)
        _set_run_font(run, font, size, bold=bold, color=color)
        return
    for index, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        if index % 2 == 1:
            _set_run_font(run, "Consolas", size - 1, color=color)
        else:
            _set_run_font(run, font, size, bold=bold, color=color)


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    _set_run_font(run, "Calibri", 9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _footer(section, title: str) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left = paragraph.add_run(f"CurioEd  ·  {title}  ·  {DATE}  ·  ")
    _set_run_font(left, "Calibri", 9, color=MUTED)
    _add_page_number(paragraph)
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("CurioEd group submission")
    _set_run_font(run, "Calibri", 9, color=MUTED)


def _add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(header)
        _set_run_font(run, "Calibri", 9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        _shade(cell, "1E3A5F")
        _set_cell_border(cell)
    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            cell = table.rows[r_index + 1].cells[c_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            _paragraph_text(paragraph, value, "Calibri", 9, color=INK)
            if r_index % 2 == 1:
                _shade(cell, "F4F7FB")
            _set_cell_border(cell)
    document.add_paragraph()


def _add_sig_table(document: Document) -> None:
    headers = ["Name", "Signature", "Date"]
    rows = [[name, "", ""] for name in MEMBERS]
    rows.append(["Additional member (name in full)", "", ""])
    table = document.add_table(rows=1 + len(rows), cols=3)
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        _set_run_font(run, "Calibri", 10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        _shade(cell, "1E3A5F")
    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            cell = table.rows[r_index + 1].cells[c_index]
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            _set_run_font(run, "Calibri", 11, color=INK)
            # Tall enough to sign.
            tc = cell._tc
            tc_pr = tc.get_or_add_tcPr()
            tc_mar = OxmlElement("w:tcMar")
            for edge, val in (("top", "160"), ("bottom", "160")):
                node = OxmlElement(f"w:{edge}")
                node.set(qn("w:w"), val)
                node.set(qn("w:type"), "dxa")
                tc_mar.append(node)
            tc_pr.append(tc_mar)


def render_docx(path: Path, title: str, blocks: list[tuple]) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    _footer(section, title)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    pf = normal.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.08

    for block in blocks:
        kind = block[0]
        if kind == "cover":
            _, heading, deck = block
            kicker = document.add_paragraph()
            kicker.paragraph_format.space_after = Pt(2)
            run = kicker.add_run(f"CURIOED  ·  GROUP SUBMISSION  ·  {DATE.upper()}")
            _set_run_font(run, "Calibri", 10, bold=True, color=MUTED)
            title_p = document.add_paragraph()
            title_p.paragraph_format.space_before = Pt(6)
            title_p.paragraph_format.space_after = Pt(6)
            run = title_p.add_run(heading)
            _set_run_font(run, "Calibri", 22, bold=True, color=NAVY)
            deck_p = document.add_paragraph()
            _paragraph_text(deck_p, deck, "Times New Roman", 12, color=INK)
            link_p = document.add_paragraph()
            run = link_p.add_run(REPO)
            _set_run_font(run, "Calibri", 10, color=NAVY)
            rule = document.add_paragraph()
            rule.paragraph_format.space_after = Pt(10)
            run = rule.add_run("─" * 42)
            _set_run_font(run, "Calibri", 10, color=NAVY)
        elif kind == "h1":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(14)
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run(block[1])
            _set_run_font(run, "Calibri", 16, bold=True, color=NAVY)
        elif kind == "h2":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(block[1])
            _set_run_font(run, "Calibri", 13, bold=True, color=NAVY)
        elif kind == "p":
            paragraph = document.add_paragraph()
            _paragraph_text(paragraph, block[1], "Times New Roman", 11, color=INK)
        elif kind == "note":
            paragraph = document.add_paragraph()
            _paragraph_text(paragraph, block[1], "Times New Roman", 11, color=INK)
        elif kind in ("ul", "ol"):
            style = "List Bullet" if kind == "ul" else "List Number"
            for item in block[1]:
                paragraph = document.add_paragraph(style=style)
                _paragraph_text(paragraph, item, "Times New Roman", 11, color=INK)
        elif kind == "table":
            _, headers, rows = block
            _add_table(document, headers, rows)
        elif kind == "sig":
            _add_sig_table(document)
        elif kind == "break":
            document.add_page_break()
        else:
            raise ValueError(kind)
    document.save(path)


def _print_pdf(html_path: Path, pdf_path: Path) -> None:
    """Headless Chrome writes the PDF and then stays up. Stop it once the file settles."""
    if pdf_path.exists():
        pdf_path.unlink()
    process = subprocess.Popen(
        [
            CHROME,
            "--headless",
            "--disable-gpu",
            "--no-sandbox",
            "--disable-dev-shm-usage",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf_path}",
            html_path.as_uri(),
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        start_new_session=True,
    )
    deadline = time.time() + 30
    last_size = -1
    stable_reads = 0
    while time.time() < deadline:
        if pdf_path.exists():
            size = pdf_path.stat().st_size
            if size > 2000 and size == last_size:
                stable_reads += 1
                if stable_reads >= 2:
                    break
            else:
                stable_reads = 0
            last_size = size
        time.sleep(0.35)
    try:
        os.killpg(process.pid, signal.SIGKILL)
    except ProcessLookupError:
        pass
    process.wait(timeout=5)
    if not pdf_path.exists() or pdf_path.stat().st_size < 2000:
        raise RuntimeError(f"Chrome did not write {pdf_path}")


def build() -> None:
    for stem, title, blocks in DOCUMENTS:
        compact = stem == "03-Data-Note"
        html_path = ROOT / f"{stem}.html"
        html_path.write_text(render_html(title, blocks, compact=compact), encoding="utf-8")
        render_docx(ROOT / f"{stem}.docx", title, blocks)
        _print_pdf(html_path, ROOT / f"{stem}.pdf")
        print(f"wrote {stem}")


if __name__ == "__main__":
    build()
